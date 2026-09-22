"""Acceptance tests for table adapters, result bundle export, and cache admin."""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ioc_rejudge.cache_admin import (
    CacheAdminError,
    apply_plan,
    build_cleanup_plan,
    cache_stats,
)
from ioc_rejudge.export_bundle import ExportConflictError, export_bundle
from ioc_rejudge.input_adapters import (
    adapt_csv,
    adapt_table,
    adapt_xlsx,
    defang,
    formula_safe_preview,
    has_formula_risk,
    restore_defang,
)


NOW = datetime(2026, 9, 22, 12, 0, tzinfo=timezone.utc)


def _verdict(ioc: str = "old.invalid") -> dict:
    return {
        "ioc": ioc,
        "original_ioc": ioc,
        "ioc_type": "domain",
        "conclusion": "灰",
        "reason": "offline test row",
        "review_suggestion": "不看",
        "confidence": "medium",
        "provider_statuses": {"test": "success"},
        "scope_actions": [],
        "retained_urls": [],
        "missing_required_providers": [],
    }


def _write_xlsx(
    path: Path,
    rows: list[tuple[object, ...]],
    *,
    sheet_name: str = "Sheet",
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _cache_shard(
    path: Path,
    *,
    result: bool,
    fetched_at: str,
    ioc: str = "old.invalid",
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if result:
        row = {
            "key": f"key-{ioc}",
            "ioc": ioc,
            "fingerprint": "fp",
            "fetched_at": fetched_at,
            "result": {"ioc": ioc},
        }
    else:
        row = {
            "key": f"key-{ioc}",
            "ioc": ioc,
            "params": {},
            "fetched_at": fetched_at,
            "raw": {"ok": True},
        }
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(row, separators=(",", ":")) + "\n")


class TestInputAdapters:
    def test_csv_reports_locations_defang_duplicates_and_invalid_rows(self, tmp_path):
        path = tmp_path / "iocs.csv"
        path.write_text(
            "indicator,ignored\n"
            "Example.INVALID,first\n"
            "example[.]invalid,second\n"
            "not a domain,invalid\n"
            ",empty\n",
            encoding="utf-8",
        )

        adapted = adapt_csv(path)

        assert adapted.report.source_type == "csv"
        assert adapted.report.selected_column == "indicator"
        assert adapted.report.total_rows == 4
        assert adapted.report.parsed_count == 1
        assert adapted.report.defang_restored_count == 1
        assert adapted.report.duplicate_count == 1
        assert adapted.report.error_count == 2
        assert adapted.bundle.kind.value == "ioc_list"
        assert [target.normalized for target in adapted.bundle.targets] == [
            "example.invalid"
        ]
        assert [row.source_line for row in adapted.report.rows] == [2, 3, 4, 5]
        assert adapted.report.rows[0].restored is False
        assert adapted.report.rows[1].restored is True
        assert adapted.report.rows[1].candidate_value == "example.invalid"
        assert adapted.report.rows[1].duplicate is True
        assert adapted.report.duplicates == [
            {
                "normalized": "example.invalid",
                "value": "Example.INVALID",
                "count": 2,
                "source_locations": ["line 2", "line 3"],
            }
        ]

    def test_csv_honors_explicit_column_map_and_gbk_fallback(self, tmp_path):
        path = tmp_path / "gbk.csv"
        path.write_bytes("alert_object\nbad.invalid\n".encode("gbk"))

        adapted = adapt_csv(path, column_map={"alert_object": "ioc"})

        assert adapted.report.selected_column == "alert_object"
        assert adapted.bundle.targets[0].normalized == "bad.invalid"

    def test_csv_quoted_newline_does_not_split_record(self, tmp_path):
        path = tmp_path / "quoted.csv"
        path.write_text(
            'ioc,note\nfirst.invalid,"multi\nline"\nsecond.invalid,ok\n',
            encoding="utf-8",
        )

        adapted = adapt_csv(path)

        assert [row.source_line for row in adapted.report.rows] == [2, 4]
        assert [row.candidate_value for row in adapted.report.rows] == [
            "first.invalid",
            "second.invalid",
        ]
        assert adapted.report.error_count == 0

    def test_formula_like_cells_are_rejected_before_parsing(self, tmp_path):
        path = tmp_path / "formulas.csv"
        path.write_text(
            "ioc\n"
            "=1+1\n"
            "+cmd\n"
            "-2\n"
            "@ATTACK\n"
            "\tTAB\n"
            "good.invalid\n",
            encoding="utf-8",
        )

        adapted = adapt_csv(path)

        assert adapted.report.formula_risk_count == 5
        assert adapted.report.error_count == 5
        assert adapted.report.parsed_count == 1
        assert adapted.bundle.targets[0].normalized == "good.invalid"
        assert all(row.formula_risk is False for row in adapted.report.rows[-1:])
        assert adapted.report.errors[0].endswith(repr("'=1+1"))

    def test_xlsx_named_sheet_header_row_and_report(self, tmp_path):
        path = tmp_path / "iocs.xlsx"
        _write_xlsx(
            path,
            [
                ("legacy", "data"),
                ("skip", "value"),
                ("metadata", "old.invalid"),
                ("metadata", "Old[.]INVALID"),
                ("metadata", "+cmd"),
            ],
            sheet_name="Alerts",
        )

        adapted = adapt_xlsx(path, sheet_name="Alerts", header_row=2)

        assert adapted.report.source_type == "xlsx"
        assert adapted.report.selected_column == "value"
        assert [row.source_line for row in adapted.report.rows] == [3, 4, 5]
        assert adapted.report.parsed_count == 1
        assert adapted.report.defang_restored_count == 1
        assert adapted.report.formula_risk_count == 1
        assert adapted.report.error_count == 1

    def test_adapt_table_dispatch_and_unsupported_suffix(self, tmp_path):
        csv_path = tmp_path / "table.csv"
        csv_path.write_text("ioc\nbad.invalid\n", encoding="utf-8")
        xlsx_path = tmp_path / "table.xlsm"
        _write_xlsx(xlsx_path, [("ioc",), ("bad.invalid",)])

        assert adapt_table(csv_path).bundle.targets[0].normalized == "bad.invalid"
        assert adapt_table(xlsx_path).bundle.targets[0].normalized == "bad.invalid"
        with pytest.raises(ValueError, match="unsupported table input type"):
            adapt_table(tmp_path / "table.tsv")

    def test_defang_helpers_round_trip_display_safe_values(self):
        assert defang("https://Example.invalid/a:8080") == (
            "hXXps[:]//Example[.]invalid/a[:]8080"
        )
        assert restore_defang("hXXp[:]//EVIL[.]invalid/a") == (
            "http://EVIL.invalid/a"
        )
        assert restore_defang("evil[DOT]invalid") == "evil.invalid"
        assert restore_defang("already.invalid") is None
        assert has_formula_risk("+cmd") is True
        assert has_formula_risk("\rCR") is True
        assert has_formula_risk("cmd") is False
        assert formula_safe_preview("@cmd") == "'@cmd"


class TestExportBundle:
    def test_writes_selected_outputs_and_optional_json_documents(self, tmp_path):
        output_dir = tmp_path / "bundle"
        diagnostics = {"parse_error_count": 0}
        diff = {"changed": 1}

        result = export_bundle(
            [_verdict()],
            output_dir=output_dir,
            base_name="results",
            jsonl=True,
            csv=True,
            xlsx=True,
            diagnostics=diagnostics,
            diff=diff,
        )

        assert result.outputs["jsonl"] == output_dir / "results.jsonl"
        assert result.outputs["csv"] == output_dir / "results.csv"
        assert result.outputs["xlsx"] == output_dir / "results.xlsx"
        assert result.outputs["diagnostics"] == output_dir / "results.diagnostics.json"
        assert result.outputs["diff"] == output_dir / "results.diff.json"
        rows = [
            json.loads(line)
            for line in result.outputs["jsonl"].read_text(encoding="utf-8").splitlines()
        ]
        assert rows[0]["ioc"] == "old.invalid"
        assert rows[0]["original_ioc"] == "old.invalid"
        assert rows[0]["conclusion"] == "灰"
        assert rows[0]["provider_statuses"] == {"test": "success"}
        with result.outputs["csv"].open(newline="", encoding="utf-8") as handle:
            assert list(csv.DictReader(handle))[0]["ioc"] == "old.invalid"
        load_workbook(result.outputs["xlsx"]).close()
        assert json.loads(result.outputs["diagnostics"].read_text("utf-8")) == diagnostics
        assert json.loads(result.outputs["diff"].read_text("utf-8")) == diff

    def test_explicit_paths_are_allowed_without_output_dir(self, tmp_path):
        paths = {
            "jsonl": tmp_path / "only.jsonl",
            "diagnostics": tmp_path / "only.diagnostics.json",
        }

        result = export_bundle(
            [_verdict()],
            output_dir=None,
            jsonl=True,
            csv=False,
            xlsx=False,
            jsonl_path=paths["jsonl"],
            diagnostics={"ok": True},
            diagnostics_path=paths["diagnostics"],
        )

        assert set(result.outputs) == {"jsonl", "diagnostics"}
        assert paths["jsonl"].exists()
        assert paths["diagnostics"].exists()

    def test_collision_preflight_writes_nothing(self, tmp_path):
        collision = tmp_path / "same.out"
        collision.write_text("KEEP", encoding="utf-8")

        with pytest.raises(ExportConflictError, match="destination collision"):
            export_bundle(
                [_verdict()],
                output_dir=tmp_path,
                jsonl_path=collision,
                csv_path=tmp_path / "." / "same.out",
                xlsx=False,
            )

        assert collision.read_text(encoding="utf-8") == "KEEP"
        assert list(tmp_path.glob("*")) == [collision]

    def test_protected_destination_preflight_writes_nothing(self, tmp_path):
        protected = tmp_path / "do-not-touch.jsonl"
        protected.write_text("KEEP", encoding="utf-8")

        with pytest.raises(ExportConflictError, match="equals protected path"):
            export_bundle(
                [_verdict()],
                output_dir=tmp_path,
                jsonl_path=protected,
                csv=False,
                xlsx=False,
                protected_paths=[protected],
            )

        assert protected.read_text(encoding="utf-8") == "KEEP"

    def test_path_inside_protected_directory_is_rejected(self, tmp_path):
        protected_dir = tmp_path / "protected"
        protected_dir.mkdir()

        with pytest.raises(ExportConflictError, match="inside protected path"):
            export_bundle(
                [_verdict()],
                output_dir=protected_dir,
                csv=False,
                xlsx=False,
                protected_paths=[protected_dir],
            )

        assert list(protected_dir.iterdir()) == []

    def test_preflight_directory_failure_does_not_write_earlier_outputs(self, tmp_path):
        locked_xlsx = tmp_path / "locked.xlsx"
        locked_xlsx.mkdir()

        with pytest.raises(OSError, match="output path is a directory"):
            export_bundle(
                [_verdict()],
                output_dir=tmp_path,
                xlsx_path=locked_xlsx,
            )

        assert not (tmp_path / "results.jsonl").exists()
        assert not (tmp_path / "results.csv").exists()
        assert locked_xlsx.is_dir()

    def test_base_name_must_be_a_simple_file_name(self, tmp_path):
        with pytest.raises(ExportConflictError, match="base_name"):
            export_bundle(
                [_verdict()],
                output_dir=tmp_path,
                base_name="../results",
                csv=False,
                xlsx=False,
            )


class TestCacheAdmin:
    def test_cache_stats_separates_result_cache_from_provider_cache(self, tmp_path):
        provider_dir = tmp_path / ".cache_test_provider"
        result_dir = tmp_path / ".cache_adjudication_results"
        _cache_shard(
            provider_dir / "cache_2026-09-01.jsonl",
            result=False,
            fetched_at="2026-09-01T00:00:00+00:00",
        )
        _cache_shard(
            result_dir / "cache_2026-09-01.jsonl",
            result=True,
            fetched_at="2026-09-01T00:00:00+00:00",
        )

        stats = cache_stats(tmp_path)

        assert stats["sections"]["provider"]["cache_directories"] == 1
        assert stats["sections"]["provider"]["files"] == 1
        assert stats["sections"]["provider"]["valid_entries"] == 1
        assert stats["sections"]["provider"]["bytes"] == (provider_dir / "cache_2026-09-01.jsonl").stat().st_size
        assert stats["sections"]["result"]["cache_directories"] == 1
        assert stats["sections"]["result"]["files"] == 1
        assert stats["sections"]["result"]["valid_entries"] == 1
        assert stats["sections"]["result"]["bytes"] == (result_dir / "cache_2026-09-01.jsonl").stat().st_size
        assert stats["total"]["valid_entries"] == 2

    def test_cache_stats_counts_malformed_rows_without_crashing(self, tmp_path):
        provider_dir = tmp_path / ".cache_test_provider"
        shard = provider_dir / "cache_2026-09-01.jsonl"
        _cache_shard(
            shard,
            result=False,
            fetched_at="2026-09-01T00:00:00+00:00",
        )
        with shard.open("a", encoding="utf-8") as handle:
            handle.write("{not-json\n")
            handle.write(json.dumps({"key": "missing-fields"}) + "\n")

        stats = cache_stats(tmp_path, cache_type="provider")

        assert stats["sections"]["provider"]["valid_entries"] == 1
        assert stats["sections"]["provider"]["invalid_entries"] == 2
        assert stats["total"]["invalid_entries"] == 2

    def test_cleanup_plan_selects_only_old_complete_shards(self, tmp_path):
        provider_dir = tmp_path / ".cache_test_provider"
        old = provider_dir / "cache_2026-09-01.jsonl"
        mixed = provider_dir / "cache_2026-09-21.jsonl"
        malformed = provider_dir / "cache_2026-09-03.jsonl"
        _cache_shard(old, result=False, fetched_at="2026-09-01T00:00:00+00:00")
        _cache_shard(
            mixed,
            result=False,
            fetched_at="2026-09-21T00:00:00+00:00",
        )
        _cache_shard(
            mixed,
            result=False,
            fetched_at="2026-09-22T00:00:00+00:00",
        )
        _cache_shard(
            malformed,
            result=False,
            fetched_at="2026-09-03T00:00:00+00:00",
        )
        with malformed.open("a", encoding="utf-8") as handle:
            handle.write("{bad\n")

        plan = build_cleanup_plan(
            tmp_path,
            cache_type="provider",
            before_date_utc=date(2026, 9, 22),
            now=NOW,
        )

        assert [item.path for item in plan.files] == [old]
        assert plan.before_date_utc == date(2026, 9, 22)
        assert plan.cutoff_utc == datetime(2026, 9, 22, tzinfo=timezone.utc)
        assert plan.eligible_bytes == old.stat().st_size

    def test_cleanup_plan_treats_naive_project_timestamps_as_utc(self, tmp_path):
        shard = tmp_path / ".cache_test_provider" / "cache_2026-09-01.jsonl"
        _cache_shard(shard, result=False, fetched_at="2026-09-01T10:00:00")

        plan = build_cleanup_plan(
            tmp_path,
            cache_type="provider",
            before_date_utc=date(2026, 9, 2),
            now=NOW,
        )

        assert [item.path for item in plan.files] == [shard]

    def test_default_apply_is_dry_run_and_leaves_files(self, tmp_path):
        old = tmp_path / ".cache_test_provider" / "cache_2026-09-01.jsonl"
        _cache_shard(old, result=False, fetched_at="2026-09-01T00:00:00+00:00")
        plan = build_cleanup_plan(
            tmp_path,
            before_date_utc=date(2026, 9, 22),
            now=NOW,
        )

        result = apply_plan(plan)

        assert result == {
            "executed": False,
            "cache_root": str(plan.cache_root),
            "would_delete_files": 1,
            "would_delete_bytes": old.stat().st_size,
        }
        assert old.exists()

    def test_apply_plan_executes_only_verified_plan(self, tmp_path):
        provider_dir = tmp_path / ".cache_test_provider"
        old = provider_dir / "cache_2026-09-01.jsonl"
        current = provider_dir / "cache_2026-09-22.jsonl"
        _cache_shard(old, result=False, fetched_at="2026-09-01T00:00:00+00:00")
        _cache_shard(
            current,
            result=False,
            fetched_at="2026-09-22T00:00:00+00:00",
        )
        plan = build_cleanup_plan(
            tmp_path,
            before_date_utc=date(2026, 9, 22),
            now=NOW,
        )
        old_size = old.stat().st_size

        result = apply_plan(plan, execute=True)

        assert result["executed"] is True
        assert result["deleted_files"] == 1
        assert result["deleted_bytes"] == old_size
        assert not old.exists()
        assert current.exists()

    def test_apply_plan_aborts_without_deletion_when_a_file_changed(self, tmp_path):
        provider_dir = tmp_path / ".cache_test_provider"
        first = provider_dir / "cache_2026-09-01.jsonl"
        second = provider_dir / "cache_2026-09-02.jsonl"
        _cache_shard(first, result=False, fetched_at="2026-09-01T00:00:00+00:00")
        _cache_shard(second, result=False, fetched_at="2026-09-02T00:00:00+00:00")
        plan = build_cleanup_plan(
            tmp_path,
            before_date_utc=date(2026, 9, 22),
            now=NOW,
        )
        with second.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps({
                "key": "new",
                "ioc": "new.invalid",
                "params": {},
                "fetched_at": "2026-09-02T01:00:00+00:00",
                "raw": {},
            }) + "\n")

        with pytest.raises(CacheAdminError, match="changed since planning"):
            apply_plan(plan, execute=True)

        assert first.exists()
        assert second.exists()

    def test_cleanup_plan_rejects_mutually_exclusive_and_unknown_options(self, tmp_path):
        with pytest.raises(CacheAdminError, match="either ttl or before_date_utc"):
            build_cleanup_plan(
                tmp_path,
                ttl=timedelta(days=1),
                before_date_utc="2026-09-01",
                now=NOW,
            )
        with pytest.raises(CacheAdminError, match="unknown cache_type"):
            cache_stats(tmp_path, cache_type="everything")

    def test_cache_admin_requires_existing_directory(self, tmp_path):
        missing = tmp_path / "missing"

        with pytest.raises(CacheAdminError, match="does not exist"):
            cache_stats(missing)
        file_root = tmp_path / "file"
        file_root.write_text("", encoding="utf-8")
        with pytest.raises(CacheAdminError, match="not a directory"):
            cache_stats(file_root)

    def test_cleanup_all_uses_result_schema_for_result_shards(self, tmp_path):
        result_shard = tmp_path / ".cache_adjudication_results" / "cache_2026-09-01.jsonl"
        _cache_shard(
            result_shard,
            result=True,
            fetched_at="2026-09-01T00:00:00+00:00",
        )

        plan = build_cleanup_plan(
            tmp_path,
            cache_type="all",
            before_date_utc=date(2026, 9, 22),
            now=NOW,
        )

        assert [item.path for item in plan.files] == [result_shard]
