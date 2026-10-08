"""Tests for jobs CLI commands, judge --queue, and offline runner."""

from __future__ import annotations

import json
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

from ioc_rejudge.job_queue import (
    DEFAULT_KEEP,
    InvalidJobStateError,
    UnifiedJobQueue,
)
from ioc_rejudge.providers.factory import DEFAULT_PROVIDERS


ROOT = Path(__file__).resolve().parents[1]

SNAPSHOT_ROW = {
    "ioc": "test-malware.invalid",
    "data": [
        {
            "key": "test-malware.invalid",
            "level": 70,
            "source": ["sample-base"],
            "family": ["trojan-downloader"],
            "hash": [
                {
                    "md5": "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaab",
                    "level": 70,
                    "time": "2026-07-15 10:00:00",
                }
            ],
            "updatetime": "2026-07-20 08:00:00",
            "context": "test-malware.invalid associated with trojan activity",
        }
    ],
}


def _run_module(*args: str, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "ioc_rejudge", *args],
        cwd=str(cwd or ROOT),
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def _queue(tmp_path: Path) -> UnifiedJobQueue:
    return UnifiedJobQueue(tmp_path / "jobs")


def test_jobs_help_lists_ten_subcommands() -> None:
    result = _run_module("jobs", "--help")
    assert result.returncode == 0, result.stderr
    text = result.stdout.lower()
    assert "list" in text
    assert "status" in text
    assert "run" in text
    assert "cancel" in text
    assert "prune" in text
    assert "results" in text
    assert "export" in text
    assert "explain" in text
    assert "review" in text
    assert "diff" in text


def _seed_succeeded_job(
    queue: UnifiedJobQueue,
    rows: list[dict],
    *,
    text: str = "alpha.invalid\nbeta.invalid\n",
) -> str:
    """Create a bare job, write results, and finish as succeeded."""
    job = queue.create_job(
        text,
        input_kind="bare",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    job_id = job["job_id"]
    claimed = queue.claim(job_id, runner_name="test", pid=1)
    assert claimed is not None
    queue.append_results(job_id, rows)
    queue.finish(
        job_id,
        state="succeeded",
        result_summary={"rows": len(rows), "conclusions": {}},
    )
    return job_id


def _sample_result_rows() -> list[dict]:
    return [
        {
            "ioc": "alpha.invalid",
            "conclusion": "误报",
            "route": "A",
            "disposition": "allow",
            "reason": "synthetic-a",
        },
        {
            "ioc": "beta.invalid",
            "conclusion": "待复核",
            "route": "B",
            "disposition": "review",
            "reason": "synthetic-b",
        },
    ]


def test_jobs_results_json_and_limit(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    rows = _sample_result_rows()
    job_id = _seed_succeeded_job(queue, rows)
    jobs_dir = str(queue.root)

    as_json = _run_module(
        "jobs", "results", job_id, "--jobs-dir", jobs_dir, "--json"
    )
    assert as_json.returncode == 0, as_json.stderr + as_json.stdout
    payload = json.loads(as_json.stdout)
    assert payload["job_id"] == job_id
    assert payload["state"] == "succeeded"
    assert len(payload["rows"]) == 2
    assert payload["rows"][0]["ioc"] == "alpha.invalid"
    assert payload["rows"][1]["ioc"] == "beta.invalid"
    assert payload.get("skipped", 0) == 0

    limited = _run_module(
        "jobs",
        "results",
        job_id,
        "--jobs-dir",
        jobs_dir,
        "--limit",
        "1",
    )
    assert limited.returncode == 0, limited.stderr + limited.stdout
    out_lines = [ln for ln in limited.stdout.splitlines() if ln.strip()]
    data_lines = [ln for ln in out_lines if "->" in ln]
    assert len(data_lines) == 1
    assert "alpha.invalid" in data_lines[0]
    assert "误报" in data_lines[0]
    assert "A/allow" in data_lines[0]

    all_rows = _run_module(
        "jobs",
        "results",
        job_id,
        "--jobs-dir",
        jobs_dir,
        "--limit",
        "0",
    )
    assert all_rows.returncode == 0, all_rows.stderr + all_rows.stdout
    data_all = [ln for ln in all_rows.stdout.splitlines() if "->" in ln]
    assert len(data_all) == 2


def test_jobs_results_non_succeeded_empty_exit_0(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "queued.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    result = _run_module(
        "jobs",
        "results",
        job["job_id"],
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert result.returncode == 0, result.stderr + result.stdout
    payload = json.loads(result.stdout)
    assert payload["state"] == "queued"
    assert payload["rows"] == []


def test_jobs_results_missing_job_exit_3(tmp_path: Path) -> None:
    jobs_dir = tmp_path / "jobs"
    jobs_dir.mkdir()
    result = _run_module(
        "jobs",
        "results",
        "jq-missing-job-00000001",
        "--jobs-dir",
        str(jobs_dir),
        "--json",
    )
    assert result.returncode == 3


def test_jobs_results_skips_bad_lines(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job_id = _seed_succeeded_job(queue, _sample_result_rows()[:1])
    results_path = queue.root / job_id / "results.jsonl"
    # Append a corrupt line after the valid one.
    with results_path.open("a", encoding="utf-8") as handle:
        handle.write("not-json\n")
        handle.write(
            json.dumps(
                {
                    "ioc": "gamma.invalid",
                    "conclusion": "灰",
                    "route": "C",
                    "disposition": "monitor",
                },
                ensure_ascii=False,
            )
            + "\n"
        )

    result = _run_module(
        "jobs", "results", job_id, "--jobs-dir", str(queue.root), "--json"
    )
    assert result.returncode == 0, result.stderr + result.stdout
    payload = json.loads(result.stdout)
    assert len(payload["rows"]) == 2
    assert payload["skipped"] == 1


def test_jobs_export_formats_and_collision_avoidance(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    rows = _sample_result_rows()
    job_id = _seed_succeeded_job(queue, rows)
    jobs_dir = str(queue.root)
    export_dir = queue.root / job_id / "export"

    csv_run = _run_module(
        "jobs",
        "export",
        job_id,
        "--format",
        "csv",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert csv_run.returncode == 0, csv_run.stderr + csv_run.stdout
    csv_payload = json.loads(csv_run.stdout)
    assert csv_payload["ok"] is True
    assert csv_payload["rows"] == 2
    csv_path = Path(csv_payload["path"])
    assert csv_path.is_file()
    assert csv_path.parent == export_dir.resolve() or csv_path.parent == export_dir
    csv_text = csv_path.read_text(encoding="utf-8")
    csv_lines = [ln for ln in csv_text.splitlines() if ln.strip()]
    assert len(csv_lines) == 3  # header + 2 data rows

    xlsx_run = _run_module(
        "jobs",
        "export",
        job_id,
        "--format",
        "xlsx",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert xlsx_run.returncode == 0, xlsx_run.stderr + xlsx_run.stdout
    xlsx_payload = json.loads(xlsx_run.stdout)
    assert xlsx_payload["ok"] is True
    assert xlsx_payload["rows"] == 2
    xlsx_path = Path(xlsx_payload["path"])
    assert xlsx_path.is_file()
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    assert "总" in wb.sheetnames
    sheet = wb["总"]
    # header + 2 data rows
    assert sheet.max_row == 3

    jsonl_run = _run_module(
        "jobs",
        "export",
        job_id,
        "--format",
        "jsonl",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert jsonl_run.returncode == 0, jsonl_run.stderr + jsonl_run.stdout
    jsonl_payload = json.loads(jsonl_run.stdout)
    assert jsonl_payload["rows"] == 2
    jsonl_path = Path(jsonl_payload["path"])
    jsonl_rows = [
        json.loads(line)
        for line in jsonl_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(jsonl_rows) == 2

    # Repeat CSV export: must not overwrite; filename increments.
    first_csv_bytes = csv_path.read_bytes()
    csv_again = _run_module(
        "jobs",
        "export",
        job_id,
        "--format",
        "csv",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert csv_again.returncode == 0, csv_again.stderr + csv_again.stdout
    again_payload = json.loads(csv_again.stdout)
    again_path = Path(again_payload["path"])
    assert again_path != csv_path
    assert again_path.is_file()
    assert csv_path.read_bytes() == first_csv_bytes
    assert again_path.name.startswith("results-")
    assert again_path.suffix == ".csv"


def test_jobs_export_rejects_non_succeeded(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "q.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    export_dir = queue.root / job["job_id"] / "export"
    before = set(export_dir.iterdir()) if export_dir.is_dir() else set()
    result = _run_module(
        "jobs",
        "export",
        job["job_id"],
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert result.returncode == 3, result.stderr + result.stdout
    after = set(export_dir.iterdir()) if export_dir.is_dir() else set()
    assert after == before


def test_jobs_export_out_conflict_refuses(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    rows = _sample_result_rows()
    job_id = _seed_succeeded_job(queue, rows)
    target = tmp_path / "custom-out.csv"
    original = b"keep-me-intact\n"
    target.write_bytes(original)

    result = _run_module(
        "jobs",
        "export",
        job_id,
        "--format",
        "csv",
        "--out",
        str(target),
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert result.returncode != 0, result.stdout
    assert target.read_bytes() == original


def test_jobs_list_and_run_print_jobs_dir(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "dir.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    jobs_dir = queue.root
    abs_dir = str(jobs_dir.resolve())

    listed = _run_module("jobs", "list", "--jobs-dir", str(jobs_dir))
    assert listed.returncode == 0, listed.stderr
    first = listed.stdout.splitlines()[0] if listed.stdout.strip() else ""
    assert first.startswith("jobs dir:")
    assert abs_dir in first or str(jobs_dir) in first

    listed_json = _run_module(
        "jobs", "list", "--jobs-dir", str(jobs_dir), "--json"
    )
    assert listed_json.returncode == 0, listed_json.stderr
    list_payload = json.loads(listed_json.stdout)
    assert "jobs_dir" in list_payload
    assert Path(list_payload["jobs_dir"]).resolve() == jobs_dir.resolve()

    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    ran = _run_module(
        "jobs",
        "run",
        job["job_id"],
        "--jobs-dir",
        str(jobs_dir),
        "--cache-dir",
        str(cache_dir),
    )
    assert ran.returncode == 0, ran.stderr + ran.stdout
    run_first = ran.stdout.splitlines()[0]
    assert run_first.startswith("jobs dir:")
    assert abs_dir in run_first or str(jobs_dir) in run_first

    ran_json = _run_module(
        "jobs",
        "run",
        "jq-missing-for-dir-check-0001",
        "--jobs-dir",
        str(jobs_dir),
        "--json",
    )
    # Missing job still includes jobs_dir in JSON when possible.
    # Use a fresh succeeded-path via list already covered; for run success JSON:
    job2 = queue.create_job(
        "dir2.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    ran_ok = _run_module(
        "jobs",
        "run",
        job2["job_id"],
        "--jobs-dir",
        str(jobs_dir),
        "--cache-dir",
        str(cache_dir),
        "--json",
    )
    assert ran_ok.returncode == 0, ran_ok.stderr + ran_ok.stdout
    run_payload = json.loads(ran_ok.stdout)
    assert "jobs_dir" in run_payload
    assert Path(run_payload["jobs_dir"]).resolve() == jobs_dir.resolve()
    # Silence unused variable if missing-job path was exploratory
    assert ran_json.returncode == 3


def test_main_module_source_routes_jobs() -> None:
    source = (ROOT / "ioc_rejudge" / "__main__.py").read_text(encoding="utf-8")
    assert 'sys.argv[1] == "jobs"' in source
    assert "jobs_cli" in source


def test_mark_cancelled_running_only(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "a.invalid\nb.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    with pytest.raises(InvalidJobStateError):
        queue.mark_cancelled(job["job_id"])

    claimed = queue.claim(job["job_id"], runner_name="t", pid=1)
    assert claimed is not None
    cancelled = queue.mark_cancelled(job["job_id"], note="pre-start cancel")
    assert cancelled["state"] == "cancelled"
    assert cancelled.get("error") == "pre-start cancel"
    assert queue.get(job["job_id"])["state"] == "cancelled"

    with pytest.raises(InvalidJobStateError):
        queue.mark_cancelled(job["job_id"])


def test_judge_queue_enqueue_list_and_run_bare(tmp_path: Path) -> None:
    jobs_dir = tmp_path / "jobs"
    cache_dir = tmp_path / "provider-cache"
    cache_dir.mkdir()

    enq = _run_module(
        "judge",
        "alpha.invalid",
        "beta.invalid",
        "--queue",
        "--jobs-dir",
        str(jobs_dir),
        "--mode",
        "offline",
    )
    assert enq.returncode == 0, enq.stderr
    assert "queued:" in enq.stdout
    job_id = enq.stdout.strip().split("queued:", 1)[1].strip().split()[0]
    assert job_id.startswith("jq-")

    listed = _run_module("jobs", "list", "--jobs-dir", str(jobs_dir), "--json")
    assert listed.returncode == 0, listed.stderr
    payload = json.loads(listed.stdout)
    jobs = payload if isinstance(payload, list) else payload.get("jobs", payload)
    if isinstance(jobs, dict) and "jobs" in jobs:
        jobs = jobs["jobs"]
    match = next(j for j in jobs if j["job_id"] == job_id)
    assert match["state"] == "queued"
    assert match["input"]["valid"] == 2

    ran = _run_module(
        "jobs",
        "run",
        job_id,
        "--jobs-dir",
        str(jobs_dir),
        "--cache-dir",
        str(cache_dir),
        "--json",
    )
    assert ran.returncode == 0, ran.stderr + ran.stdout
    run_payload = json.loads(ran.stdout)
    assert run_payload.get("state") == "succeeded" or run_payload.get("ok") is True

    job_dir = jobs_dir / job_id
    results_path = job_dir / "results.jsonl"
    assert results_path.is_file()
    rows = [
        json.loads(line)
        for line in results_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(rows) == 2
    doc = json.loads((job_dir / "job.json").read_text(encoding="utf-8"))
    assert doc["state"] == "succeeded"
    assert doc["result_summary"]["rows"] == 2
    assert isinstance(doc["result_summary"].get("conclusions"), dict)
    assert (job_dir / "diagnostics.json").is_file()


def test_judge_queue_json_summary_and_mutex(tmp_path: Path, monkeypatch) -> None:
    from ioc_rejudge import quick_cli

    jobs_dir = tmp_path / "jobs"
    called: list[list[str]] = []

    def fake_cli_main() -> None:
        called.append(list(sys.argv))
        raise SystemExit(0)

    monkeypatch.setattr("ioc_rejudge.cli.main", fake_cli_main)

    code = quick_cli.main(
        [
            "only.invalid",
            "--queue",
            "--jobs-dir",
            str(jobs_dir),
            "--json",
        ]
    )
    assert code == 0
    assert called == []  # --queue must not invoke main CLI

    # Mutual exclusion: --queue with main-CLI execution remainder is an error.
    code_bad = quick_cli.main(
        [
            "only.invalid",
            "--queue",
            "--jobs-dir",
            str(jobs_dir),
            "--jsonl",
            str(tmp_path / "out.jsonl"),
        ]
    )
    assert code_bad == 2


def test_cancel_queued_then_run_rejects(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "c.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    jobs_dir = str(queue.root)

    cancelled = _run_module(
        "jobs", "cancel", job["job_id"], "--jobs-dir", jobs_dir, "--json"
    )
    assert cancelled.returncode == 0, cancelled.stderr
    payload = json.loads(cancelled.stdout)
    assert payload.get("action") == "cancelled" or payload.get("state") == "cancelled"
    assert queue.get(job["job_id"])["state"] == "cancelled"

    ran = _run_module("jobs", "run", job["job_id"], "--jobs-dir", jobs_dir, "--json")
    assert ran.returncode == 3, ran.stderr + ran.stdout
    combined = (ran.stdout + ran.stderr).lower()
    assert "queued" in combined or "非" in (ran.stdout + ran.stderr)


def test_run_consumes_cancel_requested_before_start(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "d.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    # Simulate cancel flag set before the runner starts work (still queued).
    path = queue.root / job["job_id"] / "job.json"
    doc = json.loads(path.read_text(encoding="utf-8"))
    doc["cancel_requested"] = True
    path.write_text(
        json.dumps(doc, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )

    ran = _run_module(
        "jobs",
        "run",
        job["job_id"],
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert ran.returncode == 0, ran.stderr + ran.stdout
    final = queue.get(job["job_id"])
    assert final["state"] == "cancelled"
    results = queue.root / job["job_id"] / "results.jsonl"
    if results.exists():
        assert results.read_text(encoding="utf-8").strip() == ""


def test_run_missing_job_exits_3(tmp_path: Path) -> None:
    jobs_dir = tmp_path / "jobs"
    jobs_dir.mkdir()
    ran = _run_module(
        "jobs", "run", "jq-missing-job-00000001", "--jobs-dir", str(jobs_dir), "--json"
    )
    assert ran.returncode == 3


def test_run_non_queued_exits_3(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "e.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    claimed = queue.claim(job["job_id"], runner_name="other", pid=99)
    assert claimed is not None

    ran = _run_module(
        "jobs", "run", job["job_id"], "--jobs-dir", str(queue.root), "--json"
    )
    assert ran.returncode == 3


def test_run_jsonl_legacy_pipeline(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    text = json.dumps(SNAPSHOT_ROW, ensure_ascii=False) + "\n"
    job = queue.create_job(
        text,
        input_kind="jsonl",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="upload",
    )
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()

    ran = _run_module(
        "jobs",
        "run",
        job["job_id"],
        "--jobs-dir",
        str(queue.root),
        "--cache-dir",
        str(cache_dir),
        "--json",
    )
    assert ran.returncode == 0, ran.stderr + ran.stdout
    final = queue.get(job["job_id"])
    assert final["state"] == "succeeded"
    results_path = queue.root / job["job_id"] / "results.jsonl"
    rows = [
        json.loads(line)
        for line in results_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(rows) >= 1
    assert "conclusion" in rows[0]


def test_run_online_mode_without_credentials_succeeds(tmp_path: Path, monkeypatch) -> None:
    """Online jobs run without credentials: providers disabled, job succeeds."""
    for name in (
        "IOC_INFO_API_KEY",
        "K01_COMPROMISE_API_KEY",
        "FDP_ACCESS",
        "FDP_SECRET",
        "WHOIS_ACCESS",
        "WHOIS_SECRET",
        "PDNS_ACCESS",
        "PDNS_SECRET",
        "ICP_UC",
        "ICP_KEY",
    ):
        monkeypatch.delenv(name, raising=False)

    queue = _queue(tmp_path)
    cache_dir = tmp_path / "provider-cache"
    cache_dir.mkdir()
    job = queue.create_job(
        "online.invalid\n",
        input_kind="bare",
        mode="online",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    ran = _run_module(
        "jobs",
        "run",
        job["job_id"],
        "--jobs-dir",
        str(queue.root),
        "--cache-dir",
        str(cache_dir),
        "--json",
    )
    assert ran.returncode == 0, ran.stderr + ran.stdout
    payload = json.loads(ran.stdout)
    assert payload.get("ok") is True
    assert queue.get(job["job_id"])["state"] == "succeeded"


def test_prune_dry_run_then_apply(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    base = datetime(2026, 2, 1, tzinfo=timezone.utc)
    job_ids: list[str] = []
    for i in range(51):
        job = queue.create_job(
            f"n{i}.invalid\n",
            input_kind="bare",
            mode="offline",
            providers=["ioc_info"],
            preset="standard",
            source=f"n{i}",
        )
        path = queue.root / job["job_id"] / "job.json"
        doc = json.loads(path.read_text(encoding="utf-8"))
        stamp = (base + timedelta(seconds=i)).isoformat()
        doc["created_at"] = stamp
        doc["updated_at"] = stamp
        # Terminal-ish state for retention realism.
        doc["state"] = "succeeded"
        path.write_text(
            json.dumps(doc, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
            encoding="utf-8",
        )
        job_ids.append(job["job_id"])

    jobs_dir = str(queue.root)
    dry = _run_module(
        "jobs", "prune", "--keep", "50", "--jobs-dir", jobs_dir, "--json"
    )
    assert dry.returncode == 0, dry.stderr
    dry_payload = json.loads(dry.stdout)
    assert dry_payload["kept"] == 50
    assert len(dry_payload["removed"]) == 1
    assert dry_payload["freed_bytes"] >= 0
    assert (queue.root / job_ids[0]).is_dir()  # dry-run did not delete

    applied = _run_module(
        "jobs",
        "prune",
        "--keep",
        "50",
        "--apply",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert applied.returncode == 0, applied.stderr
    app_payload = json.loads(applied.stdout)
    assert app_payload["kept"] == 50
    assert len(app_payload["removed"]) == 1
    assert not (queue.root / job_ids[0]).exists()
    remaining = [
        p.name for p in queue.root.iterdir() if p.is_dir() and not p.name.startswith(".")
    ]
    assert len(remaining) == 50


def test_list_recovers_stale_running(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "stale.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    claimed = queue.claim(job["job_id"], runner_name="dead", pid=1)
    assert claimed is not None
    stale_ts = (datetime.now(timezone.utc) - timedelta(minutes=11)).isoformat()
    path = queue.root / job["job_id"] / "job.json"
    doc = json.loads(path.read_text(encoding="utf-8"))
    doc["runner"]["heartbeat_at"] = stale_ts
    doc["updated_at"] = stale_ts
    path.write_text(
        json.dumps(doc, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )

    listed = _run_module(
        "jobs", "list", "--jobs-dir", str(queue.root), "--json"
    )
    assert listed.returncode == 0, listed.stderr
    payload = json.loads(listed.stdout)
    jobs = payload if isinstance(payload, list) else payload.get("jobs", [])
    match = next(j for j in jobs if j["job_id"] == job["job_id"])
    assert match["state"] == "failed"
    assert match.get("error") == "runner lease expired"
    assert queue.get(job["job_id"])["state"] == "failed"


def test_secret_keys_still_rejected_via_store(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "sec.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    with pytest.raises((InvalidJobStateError, ValueError)):
        queue.append_results(job["job_id"], [{"ioc": "x", "api_key": "nope"}])
    with pytest.raises((InvalidJobStateError, ValueError)):
        queue.write_diagnostics(job["job_id"], {"token": "secret"})


def test_status_command(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job = queue.create_job(
        "status.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    result = _run_module(
        "jobs", "status", job["job_id"], "--jobs-dir", str(queue.root), "--json"
    )
    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["job_id"] == job["job_id"]
    assert payload["state"] == "queued"


def test_jobs_explain_second_row_and_missing(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    rows = _sample_result_rows()
    job_id = _seed_succeeded_job(queue, rows)
    jobs_dir = str(queue.root)
    second_id = f"{job_id}-000002"

    explained = _run_module(
        "jobs",
        "explain",
        job_id,
        "--result-id",
        second_id,
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert explained.returncode == 0, explained.stderr + explained.stdout
    payload = json.loads(explained.stdout)
    assert payload["ioc"] == "beta.invalid"
    assert payload["result_id"] == second_id
    assert payload["conclusion"] == "待复核"
    assert "evidence_fingerprint" in payload

    missing = _run_module(
        "jobs",
        "explain",
        job_id,
        "--result-id",
        f"{job_id}-009999",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert missing.returncode == 3


def test_jobs_review_idempotent_and_explain_shows_label(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    rows = _sample_result_rows()
    job_id = _seed_succeeded_job(queue, rows)
    jobs_dir = str(queue.root)
    results_path = queue.root / job_id / "results.jsonl"
    before = results_path.read_bytes()

    for _ in range(2):
        reviewed = _run_module(
            "jobs",
            "review",
            job_id,
            "--ioc",
            "alpha.invalid",
            "--label",
            "approved",
            "--note",
            "ops-check",
            "--reviewer",
            "analyst",
            "--jobs-dir",
            jobs_dir,
            "--json",
        )
        assert reviewed.returncode == 0, reviewed.stderr + reviewed.stdout
        body = json.loads(reviewed.stdout)
        assert body.get("ok") is True or body.get("label") == "approved"

    review_path = queue.root / job_id / "review.jsonl"
    records = [
        json.loads(line)
        for line in review_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(records) == 2
    assert results_path.read_bytes() == before

    explained = _run_module(
        "jobs",
        "explain",
        job_id,
        "--result-id",
        f"{job_id}-000001",
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert explained.returncode == 0, explained.stderr + explained.stdout
    payload = json.loads(explained.stdout)
    assert payload["conclusion"] == "误报"
    review = payload.get("review")
    assert isinstance(review, dict)
    assert review.get("label") == "approved"


def test_jobs_review_invalid_label_exit_2(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job_id = _seed_succeeded_job(queue, _sample_result_rows())
    result = _run_module(
        "jobs",
        "review",
        job_id,
        "--ioc",
        "alpha.invalid",
        "--label",
        "totally-invalid-label",
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert result.returncode == 2


def test_jobs_diff_summary_and_json(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    baseline_rows = [
        {
            "ioc": "alpha.invalid",
            "conclusion": "存活有效",
            "disposition": "block",
            "reason": "old",
        },
        {
            "ioc": "beta.invalid",
            "conclusion": "误报",
            "disposition": "allow",
            "reason": "stable",
        },
    ]
    current_rows = [
        {
            "ioc": "alpha.invalid",
            "conclusion": "误报",
            "disposition": "allow",
            "reason": "flipped",
        },
        {
            "ioc": "beta.invalid",
            "conclusion": "误报",
            "disposition": "allow",
            "reason": "stable",
            "review_suggestion": "必看",
        },
    ]
    baseline_id = _seed_succeeded_job(queue, baseline_rows)
    current_id = _seed_succeeded_job(queue, current_rows)
    jobs_dir = str(queue.root)

    human = _run_module(
        "jobs",
        "diff",
        current_id,
        "--baseline",
        baseline_id,
        "--jobs-dir",
        jobs_dir,
    )
    assert human.returncode == 0, human.stderr + human.stdout
    text = human.stdout
    assert "operations=" in text
    assert "black_to_white=" in text or "changed=" in text

    as_json = _run_module(
        "jobs",
        "diff",
        current_id,
        "--baseline",
        baseline_id,
        "--jobs-dir",
        jobs_dir,
        "--json",
    )
    assert as_json.returncode == 0, as_json.stderr + as_json.stdout
    payload = json.loads(as_json.stdout)
    body = payload.get("diff", payload)
    assert body["operations"] == 2
    assert "operational_changes" in body
    assert any(item["ioc"] == "beta.invalid" for item in body["operational_changes"])


def test_jobs_diff_rejects_queued_baseline(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    current_id = _seed_succeeded_job(queue, _sample_result_rows())
    queued = queue.create_job(
        "queued.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    result = _run_module(
        "jobs",
        "diff",
        current_id,
        "--baseline",
        queued["job_id"],
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert result.returncode == 3


def test_jobs_explain_reports_skipped_bad_lines(tmp_path: Path) -> None:
    queue = _queue(tmp_path)
    job_id = _seed_succeeded_job(queue, _sample_result_rows()[:1])
    path = queue.root / job_id / "results.jsonl"
    with path.open("a", encoding="utf-8") as handle:
        handle.write("broken-line\n")
        handle.write(
            json.dumps(
                {
                    "ioc": "beta.invalid",
                    "conclusion": "待复核",
                    "route": "B",
                    "disposition": "review",
                    "reason": "after-bad",
                },
                ensure_ascii=False,
            )
            + "\n"
        )
    # Second valid row is index 2 among valid rows after skip.
    result = _run_module(
        "jobs",
        "explain",
        job_id,
        "--result-id",
        f"{job_id}-000002",
        "--jobs-dir",
        str(queue.root),
        "--json",
    )
    assert result.returncode == 0, result.stderr + result.stdout
    payload = json.loads(result.stdout)
    assert payload["ioc"] == "beta.invalid"
    assert payload.get("skipped", 0) == 1


def test_progress_tee_isolates_handler_errors() -> None:
    from ioc_rejudge.jobs_cli import _progress_tee

    seen: list[str] = []

    def bad(*_a, **_k):
        raise RuntimeError("render boom")

    def good(*args, **_k):
        seen.append(str(args[0]) if args else "")

    tee = _progress_tee([bad, good, bad])
    tee("provider 'whois': completed in 0.1s (1 target(s))")
    assert seen == ["provider 'whois': completed in 0.1s (1 target(s))"]


def test_jobs_run_human_progress_non_tty_keeps_heartbeat(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    from ioc_rejudge.jobs_cli import main
    from ioc_rejudge.progress import LiveProgress

    queue = _queue(tmp_path)
    cache_dir = tmp_path / "provider-cache"
    cache_dir.mkdir()
    job = queue.create_job(
        "prog-a.invalid\nprog-b.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=list(DEFAULT_PROVIDERS),
        preset="standard",
        source="test",
    )
    job_id = job["job_id"]
    started = queue.get(job_id)["updated_at"]

    beats = {"n": 0}
    real_hb = UnifiedJobQueue.heartbeat

    def counting_hb(self, jid, *args, **kwargs):
        beats["n"] += 1
        return real_hb(self, jid, *args, **kwargs)

    monkeypatch.setattr(UnifiedJobQueue, "heartbeat", counting_hb)

    # Force non-TTY plain/throttled rendering regardless of capture environment.
    real_init = LiveProgress.__init__

    def plain_init(self, stream=None, *, tty=None):
        real_init(self, stream=stream, tty=False)

    monkeypatch.setattr(LiveProgress, "__init__", plain_init)

    code = main(
        [
            "run",
            job_id,
            "--jobs-dir",
            str(queue.root),
            "--cache-dir",
            str(cache_dir),
        ]
    )
    assert code == 0
    captured = capsys.readouterr()
    combined = captured.out + "\n" + captured.err
    assert f"succeeded: {job_id}" in captured.out
    assert "provider '" in combined
    assert beats["n"] >= 1
    finished = queue.get(job_id)
    assert finished["state"] == "succeeded"
    assert finished.get("updated_at")
    assert finished["updated_at"] >= started


def test_jobs_run_json_stdout_is_single_object_without_progress(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    from ioc_rejudge.jobs_cli import main
    from ioc_rejudge.progress import LiveProgress

    created = {"n": 0}
    real_init = LiveProgress.__init__

    def tracking_init(self, *args, **kwargs):
        created["n"] += 1
        return real_init(self, *args, **kwargs)

    monkeypatch.setattr(LiveProgress, "__init__", tracking_init)

    queue = _queue(tmp_path)
    cache_dir = tmp_path / "provider-cache"
    cache_dir.mkdir()
    job = queue.create_job(
        "json-prog.invalid\n",
        input_kind="bare",
        mode="offline",
        providers=["ioc_info"],
        preset="standard",
        source="test",
    )
    code = main(
        [
            "run",
            job["job_id"],
            "--jobs-dir",
            str(queue.root),
            "--cache-dir",
            str(cache_dir),
            "--json",
        ]
    )
    assert code == 0
    captured = capsys.readouterr()
    lines = [ln for ln in captured.out.splitlines() if ln.strip()]
    assert len(lines) == 1
    payload = json.loads(lines[0])
    assert payload.get("ok") is True
    assert payload.get("job_id") == job["job_id"]
    assert created["n"] == 0
    assert "provider '" not in captured.out
